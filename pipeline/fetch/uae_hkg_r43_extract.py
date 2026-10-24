# /// script
# dependencies = ["openpyxl>=3.1"]
# ///
"""Extract the small tables kept from the round-43 browser fetch (pipeline/fetch/uae_hkg_r43.mjs).

Reads research/2026-10/raw/uae_hkg_r43/ (the downloads, treated as untrusted data: openpyxl in read-only, data-only
mode; the zip is read with zipfile, nothing is executed) and writes, in the same folder:
  scad_abu_dhabi_russia.csv   Abu Dhabi ports, non-oil imports from and re-exports to Russia, AED million, per release;
                              where Russia is not in the top-10 table, the value is empty and `upper_bound` holds the
                              10th country's value
  fca_uae_russia_2010_2019.csv  Federal Customs Authority, UAE imports / exports / re-exports with Russia, AED and USD
  comtrade_hkg_7108_monthly.csv  Hong Kong imports from Russia, HS 7108 and TOTAL, USD, per month
"""

import csv
import io
import json
import re
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import openpyxl

RAW = Path(__file__).resolve().parents[2] / "research/2026-10/raw/uae_hkg_r43"
AED_PER_USD = 3.6725  # the dirham's fixed peg
SKIP = {"TOTAL", "OTHERS", "OTHER", "COUNTRY"}


def accessed(f: Path) -> str:
    return datetime.fromtimestamp(f.stat().st_mtime, tz=UTC).date().isoformat()


def top10(ws) -> tuple[dict[str, float], float]:
    """Country -> value (the year total for annual files, the month for monthly files) and the smallest listed value."""
    rows = {}
    for r in ws.iter_rows(values_only=True):
        v = [c for c in r if c is not None]
        nums = [x for x in v if isinstance(x, (int, float))]
        if v and isinstance(v[0], str) and nums and v[0].strip().upper() not in SKIP:
            rows[v[0].strip().upper()] = float(nums[-1])
    return rows, min(rows.values())


def scad() -> list[dict]:
    log = {}
    for f in RAW.glob("fetch_log_*.json"):
        for rec in json.loads(f.read_text()):
            if rec.get("file"):
                log[rec["file"]] = rec
    out = []
    for f in sorted(RAW.glob("scad_*.xlsx")):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        rec = log.get(f.name, {})
        for table, flow in (("Table 8", "imports_from_russia"), ("Table 7", "reexports_to_russia")):
            rows, low = top10(wb[table])
            val = rows.get("RUSSIA")
            out.append(
                {
                    "release": f.stem.removeprefix("scad_"),
                    "flow": flow,
                    "aed_million": round(val, 1) if val is not None else "",
                    "upper_bound_aed_million": "" if val is not None else round(low, 1),
                    "usd_million": round(val / AED_PER_USD, 1) if val is not None else "",
                    "source": rec.get("url", ""),
                    "page": rec.get("page", ""),
                    "accessed": (rec.get("at") or "")[:10],
                }
            )
    return out


def fca() -> list[dict]:
    z = RAW / "bayanat_fca_by_country.zip"
    if not z.exists():
        return []
    with zipfile.ZipFile(z) as zf:
        name = next(n for n in zf.namelist() if n.endswith(".xlsx") and "metadata" not in n.lower())
        wb = openpyxl.load_workbook(io.BytesIO(zf.read(name)), read_only=True, data_only=True)
    out = []
    for r in wb.worksheets[0].iter_rows(values_only=True):
        if len(r) >= 10 and isinstance(r[2], str) and r[2].strip() == "Russian Federation":
            imp, exp, rex = (float(r[i] or 0) for i in (5, 7, 9))
            out.append(
                {
                    "year": r[0],
                    "imports_aed": round(imp),
                    "exports_aed": round(exp),
                    "reexports_aed": round(rex),
                    "imports_usd_bn": round(imp / AED_PER_USD / 1e9, 3),
                    "exports_usd_bn": round(exp / AED_PER_USD / 1e9, 3),
                    "reexports_usd_bn": round(rex / AED_PER_USD / 1e9, 3),
                    "source": "https://bayanat.ae/en/Datasets (Federal Customs Authority, General Total Trade volume by country 2010-2019)",
                    "accessed": accessed(z),
                }
            )
    return out


def comtrade() -> list[dict]:
    out = []
    for f in sorted(RAW.glob("comtrade_hkg_7108_m_*.json")):
        period = re.search(r"(\d{6})", f.name).group(1)
        vals = {r["cmdCode"]: r["primaryValue"] for r in json.loads(f.read_text()).get("data") or []}
        if vals:
            out.append(
                {
                    "period": period,
                    "hs7108_usd_m": round(vals.get("7108", 0) / 1e6, 1),
                    "total_usd_m": round(vals.get("TOTAL", 0) / 1e6, 1),
                    "accessed": accessed(f),
                    "source": f"https://comtradeapi.un.org/public/v1/preview/C/M/HS?reporterCode=344&partnerCode=643&period={period}&flowCode=M&cmdCode=7108,TOTAL",
                }
            )
    return out


def write(name: str, rows: list[dict]) -> None:
    if not rows:
        print(f"{name}: nothing to write", file=sys.stderr)
        return
    with open(RAW / name, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    print(f"{name}: {len(rows)} rows")


if __name__ == "__main__":
    write("scad_abu_dhabi_russia.csv", scad())
    write("fca_uae_russia_2010_2019.csv", fca())
    write("comtrade_hkg_7108_monthly.csv", comtrade())
